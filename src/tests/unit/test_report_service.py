from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook

from src.reporting.report_service import generate_daily_reports


class _Database:
    def fetchone(self, _query, _parameters=()):
        return {"count": 4}


class _Repository:
    def __init__(self, now):
        self.events = [
            {
                "occurred_at": (now - timedelta(hours=2)).isoformat(),
                "severity": "high",
                "source": "process",
                "event_type": "suspicious_process",
                "description": "<script>needs review</script>",
                "risk_score": 72.0,
            },
            {
                "occurred_at": (now - timedelta(days=2)).isoformat(),
                "severity": "critical",
                "source": "old",
                "event_type": "old_event",
                "description": "Outside report period",
            },
        ]
        self.alerts = [
            {
                "created_at": (now - timedelta(hours=1)).isoformat(),
                "severity": "high",
                "title": "Rule match",
                "source": "process",
                "status": "open",
                "rationale": "Requires review",
            },
            {
                "created_at": (now - timedelta(hours=1)).isoformat(),
                "severity": "critical",
                "title": "Resolved rule",
                "source": "process",
                "status": "resolved",
            },
        ]

    def list_events(self, limit=500):
        assert limit == 10000
        return self.events

    def list_alerts(self, limit=500, status=""):
        assert limit == 10000
        return self.alerts


def test_daily_reports_include_only_recent_records_and_supported_formats(tmp_path):
    now = datetime(2026, 10, 13, 12, tzinfo=UTC)
    reports = generate_daily_reports(tmp_path, _Database(), _Repository(now), now)

    assert set(reports) == {"pdf", "xlsx", "html"}
    assert all(path.is_file() for path in reports.values())
    assert reports["pdf"].read_bytes().startswith(b"%PDF")
    html = reports["html"].read_text(encoding="utf-8")
    assert "needs review" in html
    assert "&lt;script&gt;" in html
    assert "Outside report period" not in html
    assert "Rule alerts require analyst review" in html
    workbook = load_workbook(reports["xlsx"], read_only=True)
    assert workbook.sheetnames == [
        "Executive summary",
        "Security events",
        "Active alerts",
    ]
    assert workbook["Executive summary"]["B4"].value == 1
    assert workbook["Executive summary"]["B8"].value == 4
