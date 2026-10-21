"""Excel export for the local daily executive summary."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.reporting.executive_summary import DailyExecutiveSummary


def export_xlsx(summary: DailyExecutiveSummary, destination: str | Path) -> Path:
    path = Path(destination)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Executive summary"
    overview.append(("Metric", "Value"))
    overview.append(("Generated at (UTC)", summary.generated_at.isoformat()))
    overview.append(("Period starts (UTC)", summary.period_start.isoformat()))
    overview.append(("Recorded events", len(summary.events)))
    overview.append(("Active rule alerts", len(summary.alerts)))
    overview.append(("Critical events", summary.critical_severity_count))
    overview.append(("High-severity events", summary.high_severity_count))
    overview.append(("Baseline files (not malware scans)", summary.baseline_file_count))

    events = workbook.create_sheet("Security events")
    events.append(("Time", "Severity", "Source", "Type", "Details", "Risk score"))
    for event in summary.events:
        events.append(
            (
                str(event.get("occurred_at", "")),
                str(event.get("severity", "")).upper(),
                str(event.get("source", "")),
                str(event.get("event_type", "")),
                str(event.get("description", "")),
                event.get("risk_score", ""),
            )
        )
    alerts = workbook.create_sheet("Active alerts")
    alerts.append(("Created", "Severity", "Alert", "Source", "Status", "Rationale"))
    for alert in summary.alerts:
        alerts.append(
            (
                str(alert.get("created_at", "")),
                str(alert.get("severity", "")).upper(),
                str(alert.get("title", "")),
                str(alert.get("source", "")),
                str(alert.get("status", "")),
                str(alert.get("rationale", "")),
            )
        )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123C64")
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column]
            width = min(max(max(values, default=0) + 2, 12), 72)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    workbook.save(path)
    return path
